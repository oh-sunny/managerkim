"""Build deterministic, fictional Korean HR notice test fixtures.

Requires python-docx and reportlab. Run from the repository root.
"""

from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.utils import simpleSplit
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, KeepTogether


ROOT = Path(__file__).resolve().parent
DOC_DIR = ROOT / "documents"
TEXT_DIR = ROOT / "source_text"
DATA_DIR = ROOT / "data"
for folder in (DOC_DIR, TEXT_DIR, DATA_DIR):
    folder.mkdir(parents=True, exist_ok=True)

FONT = "C:/Windows/Fonts/malgun.ttf"
pdfmetrics.registerFont(TTFont("Malgun", FONT))


CASES = [
    {
        "id": "health-2026", "name": "2026 임직원 건강검진", "type": "복지",
        "target": "재직 중인 검진 대상자 120명", "deadlineAt": "2026-10-23T18:00",
        "eventDate": "2026-11-30", "confirmationMode": "separate",
        "applicationUrl": "https://intra.gaonlab.example/welfare/health/2026/apply",
        "owner": "피플운영팀 박서연", "targetCount": 120, "requiredCount": 30,
        "appliedCount": 108, "requiredAppliedCount": 25,
        "noticeMustInclude": ["검진 기관과 희망 일정을 선택", "예약 확정 안내"],
        "noticeMustAvoid": ["검진 예약 즉시 확정", "2025년 신청"],
        "documents": ["health-policy", "health-booking", "health-update"],
    },
    {
        "id": "learning-2026q4", "name": "2026년 4분기 자기계발비 지원", "type": "복지",
        "target": "재직 임직원 중 신청 요건을 충족하는 80명", "deadlineAt": "2026-11-06T18:00",
        "eventDate": "2026-12-31", "confirmationMode": "separate",
        "applicationUrl": "https://intra.gaonlab.example/welfare/learning/2026-q4/apply",
        "owner": "피플운영팀 김유진", "targetCount": 80, "requiredCount": 0,
        "appliedCount": 30, "requiredAppliedCount": 0,
        "noticeMustInclude": ["교육 과정과 예상 비용", "검토 후 결과"],
        "noticeMustAvoid": ["선착순 자동 승인", "무조건 지원"],
        "documents": ["learning-policy", "learning-review", "learning-faq"],
    },
    {
        "id": "workshop-2026", "name": "2026 가을 워크숍", "type": "행사",
        "target": "워크숍 초청 동료 100명", "deadlineAt": "2026-10-20T18:00",
        "eventDate": "2026-11-06", "confirmationMode": "separate",
        "applicationUrl": "https://intra.gaonlab.example/events/autumn-workshop-2026/rsvp",
        "owner": "컬처팀 이지훈", "targetCount": 100, "requiredCount": 12,
        "appliedCount": 51, "requiredAppliedCount": 9,
        "noticeMustInclude": ["참석 가능 여부", "좌석 확정 안내"],
        "noticeMustAvoid": ["10월 18일 마감", "신청 즉시 참석 확정"],
        "obsoleteDeadlines": ["2026-10-18"],
        "obsoleteUrls": ["https://intra.gaonlab.example/events/autumn-workshop-2026/old-rsvp"],
        "documents": ["workshop-plan", "workshop-logistics", "workshop-rsvp-update"],
    },
    {
        "id": "photo-2026", "name": "2026 사내 사진 공모전", "type": "이벤트",
        "target": "재직 임직원 80명", "deadlineAt": "2026-10-30T18:00",
        "eventDate": "2026-11-12", "confirmationMode": "immediate",
        "applicationUrl": "https://intra.gaonlab.example/events/photo-2026/submit",
        "owner": "컬처팀 최민지", "targetCount": 80, "requiredCount": 0,
        "appliedCount": 0, "requiredAppliedCount": 0,
        "noticeMustInclude": ["사진 1점과 설명", "제출 완료 화면"],
        "noticeMustAvoid": ["참여 필수", "모든 사진을 사외 홍보에 사용"],
        "documents": ["photo-plan", "photo-submit", "photo-rights"],
    },
]


def doc(id, case, ext, title, date, version, purpose, sections, status="확정"):
    return dict(id=id, case=case, ext=ext, title=title, date=date, version=version,
                status=status, purpose=purpose, sections=sections)


DOCS = [
    doc("health-policy", "health-2026", "pdf", "2026 임직원 건강검진 운영 기준", "2026-10-02", "2.0", "대상과 검진 범위의 기준을 확정한다", [
        ("운영 개요", ["올해 건강검진은 대상 동료가 검진 기관과 희망 일정을 직접 선택하는 방식으로 운영한다. 검진은 2026년 11월 30일까지 완료해야 한다.", "대상은 재직 중인 검진 대상자 120명이다. 법정 검진 일정 관리가 필요한 30명은 필수 신청 대상이며, 나머지 90명은 자율 신청 대상이다."]),
        ("신청과 확정", ["신청 마감: 2026년 10월 23일 오후 6시", "신청 화면에서 검진 기관과 희망 일정을 선택한다. 제출 후 예약 가능 여부를 확인하고 별도의 예약 확정 안내를 보낸다. 신청 완료와 검진 예약 확정은 다르다."]),
        ("유의 사항", ["검진 기관별 예약 가능 일정은 다를 수 있다. 원하는 날짜가 보이지 않으면 다른 날짜를 선택하거나 피플운영팀에 문의한다.", "민감한 건강 정보는 공지나 팀 채널에서 수집하지 않는다."]),
    ]),
    doc("health-booking", "health-2026", "docx", "건강검진 예약 신청 절차와 문의 대응", "2026-10-02", "1.3", "직원이 신청 화면에서 해야 할 일을 안내한다", [
        ("신청 절차", ["신청 링크: https://intra.gaonlab.example/welfare/health/2026/apply", "사내 계정으로 로그인한 뒤 검진 기관과 희망 일정을 선택하고 제출한다. 화면에 신청 접수 번호가 나타나면 접수가 완료된 것이다.", "예약 확정 안내는 신청 현황을 확인한 뒤 별도로 전달한다. 접수 번호만으로 검진 기관 방문 일정이 확정된 것은 아니다."]),
        ("문의와 변경", ["로그인이 되지 않거나 예약 가능 일정이 보이지 않으면 피플운영팀 박서연에게 문의한다.", "신청 뒤 일정 변경이 필요하면 신청 화면의 변경 메뉴를 사용한다. 마감 뒤 변경은 담당자에게 먼저 문의한다."]),
    ]),
    doc("health-update", "health-2026", "pdf", "건강검진 신청 일정 확정 안내", "2026-10-05", "1.0", "올해 적용할 마감과 운영 종료일을 재확인한다", [
        ("확정 일정", ["신청 마감은 2026년 10월 23일 오후 6시이고 검진 이용 종료일은 2026년 11월 30일이다.", "전년도 안내의 일정과 링크를 재사용하지 않는다. 올해 신청 링크는 예약 신청 절차 문서를 따른다."]),
        ("안내 대상", ["필수 신청 대상자에게는 미신청 상태를 확인해 개별 안내할 수 있다. 자율 신청 대상자는 참여를 강요하는 표현을 사용하지 않는다."]),
    ]),
    doc("learning-policy", "learning-2026q4", "pdf", "2026년 4분기 자기계발비 지원 기준", "2026-09-29", "1.0", "지원 범위와 신청 자격을 정한다", [
        ("제도 목적과 대상", ["직무 역량 개발에 필요한 외부 교육 과정과 도서 구입을 지원한다. 4분기 신청 대상은 재직 임직원 중 요건을 충족하는 80명이다. 신청은 자율이다.", "1인당 분기 지원 한도는 20만 원이다. 실제 지급액은 증빙과 기준 적합성 검토 뒤 확정한다."]),
        ("지원 가능 항목", ["업무 연관성이 확인되는 교육 과정, 세미나, 직무 도서를 신청할 수 있다. 개인 취미 목적의 구독료와 이미 회사가 별도로 제공하는 교육은 제외한다.", "지원 대상 여부가 불분명한 비용은 결제 전에 피플운영팀에 문의한다."]),
        ("신청 일정", ["신청 마감: 2026년 11월 6일 오후 6시", "승인 뒤 2026년 12월 31일까지 교육 수강이나 도서 구매를 완료한다."]),
    ]),
    doc("learning-review", "learning-2026q4", "docx", "자기계발비 신청 검토 절차", "2026-09-30", "1.1", "운영자가 접수 뒤 확인할 항목을 기록한다", [
        ("직원이 제출할 내용", ["신청 링크: https://intra.gaonlab.example/welfare/learning/2026-q4/apply", "신청 화면에 교육 과정 또는 도서명, 예상 비용, 직무 관련성을 입력한다. 필요한 경우 교육 소개 자료를 첨부한다."]),
        ("검토 절차", ["피플운영팀은 신청 요건과 예산을 확인한 뒤 결과를 별도로 안내한다. 신청만으로 지원이 확정되지 않는다.", "승인된 건은 결제 증빙을 제출한 뒤 정산한다. 운영자의 내부 검토 항목과 예산 잔액은 전사 공지에 포함하지 않는다."]),
    ]),
    doc("learning-faq", "learning-2026q4", "pdf", "자기계발비 직원 문의 답변", "2026-10-01", "1.0", "신청 전에 자주 묻는 질문을 정리한다", [
        ("신청 관련", ["여러 과정을 신청할 수 있으나 분기 합계 지원 한도는 20만 원이다. 신청 기한은 2026년 11월 6일 오후 6시다.", "접수 후에는 검토 결과가 별도로 안내된다. 선착순 자동 승인 방식은 아니다."]),
        ("변경과 문의", ["과정이 변경되면 기존 신청 내역을 수정하고 피플운영팀 김유진에게 알린다.", "문서에 적힌 금액은 최대 지원 한도이며 개인별 지급 약속이 아니다."]),
    ]),
    doc("workshop-plan", "workshop-2026", "docx", "2026 가을 워크숍 기획안", "2026-09-25", "0.8", "초기 운영안을 공유한다", [
        ("기획 배경", ["팀 간 협업 사례를 나누고 다음 분기 공동 과제를 논의하기 위해 2026년 11월 6일 가을 워크숍을 연다. 초청 동료는 100명이며 운영 진행 역할의 12명은 참석 여부 회신이 필요하다."]),
        ("초기 모집안", ["초기안 신청 마감: 2026년 10월 18일 오후 6시", "초기안 신청 링크: https://intra.gaonlab.example/events/autumn-workshop-2026/old-rsvp", "이 문서는 확정 전 초안이다. 모집 마감과 링크는 이후 확정 안내를 기준으로 다시 확인한다."]),
        ("참석 확정", ["행사장 좌석을 확인한 뒤 참석 확정 안내를 발송한다. 회신 완료와 좌석 확정은 다른 상태다."]),
    ], status="초안"),
    doc("workshop-logistics", "workshop-2026", "pdf", "가을 워크숍 장소와 이동 안내", "2026-10-01", "1.0", "참석자에게 필요한 현장 정보를 준비한다", [
        ("행사 일정", ["행사일: 2026년 11월 6일", "장소는 가온랩 서초 교육장이다. 세부 세션 시간표와 이동 안내는 좌석 확정자에게 별도로 제공한다."]),
        ("현장 준비", ["점심 식사는 제공한다. 식이 제한이나 이동 지원이 필요하면 참석 회신 화면의 요청란에 적거나 컬처팀 이지훈에게 문의한다.", "개별 요청 사유는 공개 채널 공지에 옮기지 않는다."]),
    ]),
    doc("workshop-rsvp-update", "workshop-2026", "docx", "가을 워크숍 참석 회신 확정 안내", "2026-10-03", "1.0", "최종 회신 기한과 링크를 확정한다", [
        ("최종 회신 방법", ["신청 마감: 2026년 10월 20일 오후 6시", "신청 링크: https://intra.gaonlab.example/events/autumn-workshop-2026/rsvp", "초청받은 동료는 링크에서 참석 가능 여부를 회신한다. 참석이 어려운 경우에도 불참을 선택해 회신할 수 있다."]),
        ("확정과 안내", ["좌석과 진행 역할을 확인한 뒤 참석 확정 안내를 별도로 보낸다. 회신 직후 참석 확정으로 표시하지 않는다.", "9월 25일 기획안의 10월 18일 기한과 이전 링크는 사용하지 않는다."]),
    ]),
    doc("photo-plan", "photo-2026", "docx", "2026 사내 사진 공모전 운영안", "2026-10-02", "1.0", "공모 주제와 참여 일정을 정한다", [
        ("공모 개요", ["일상 속 협업과 일하는 공간을 주제로 사진을 모집한다. 재직 임직원 80명이 자율적으로 참여할 수 있다.", "작품 제출 마감은 2026년 10월 30일 오후 6시이고 결과 발표는 2026년 11월 12일이다."]),
        ("제출 기준", ["참가자당 사진 1점과 100자 이내 설명을 제출한다. 타인의 얼굴이 식별되는 사진은 촬영 및 제출에 필요한 동의를 먼저 확인해야 한다.", "제출을 완료한 뒤 화면에 접수 완료 표시가 나타나는지 확인한다."]),
    ]),
    doc("photo-submit", "photo-2026", "pdf", "사진 공모전 제출 방법과 자주 묻는 질문", "2026-10-03", "1.0", "참가자가 제출 화면에서 따라 할 절차를 안내한다", [
        ("온라인 제출", ["신청 링크: https://intra.gaonlab.example/events/photo-2026/submit", "사내 계정으로 로그인하고 사진 1점과 설명을 입력해 제출한다. 제출 완료 화면이 나타나면 접수된 것이다."]),
        ("수정과 문의", ["마감 전까지 제출 화면에서 작품을 바꿀 수 있다. 파일 업로드 오류가 나면 컬처팀 최민지에게 문의한다.", "메신저로 사진만 보내면 공모전 접수가 완료되지 않는다."]),
    ]),
    doc("photo-rights", "photo-2026", "docx", "사진 공모전 심사와 활용 기준", "2026-10-04", "1.0", "제출 작품의 심사와 활용 범위를 확정한다", [
        ("심사와 발표", ["운영팀은 주제 적합성과 사진 설명을 검토해 2026년 11월 12일 결과를 발표한다. 참가자별 평가 메모는 공개하지 않는다."]),
        ("작품 활용", ["제출 작품은 사내 전시와 사내 소식에 사용될 수 있다. 사외 홍보에 쓰려면 별도 동의를 받는다.", "공지에는 참여가 자율이라는 점과 사람의 얼굴이 식별되는 사진의 동의 필요성을 명확히 적는다."]),
    ]),
]


def source_text(item):
    lines = [item["title"], f"문서번호: GL-{item['id'].upper()}", f"버전: {item['version']}",
             f"작성일: {item['date']}", f"상태: {item['status']}", f"목적: {item['purpose']}"]
    for heading, paragraphs in item["sections"]:
        lines += ["", heading, *paragraphs]
    return "\n".join(lines) + "\n"


def make_docx(item, target):
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.1)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.4)
    styles = document.styles
    for name in ("Normal", "Title", "Heading 1"):
        styles[name].font.name = "Malgun Gothic"
        styles[name].font.color.rgb = RGBColor(0, 0, 0)
    styles["Normal"].font.size = Pt(9.5)
    styles["Normal"].paragraph_format.space_after = Pt(7)
    styles["Title"].font.size = Pt(19)
    title_ppr = styles["Title"].element.get_or_add_pPr()
    for border in title_ppr.findall(qn("w:pBdr")):
        title_ppr.remove(border)
    styles["Subtitle"].font.name = "Malgun Gothic"
    styles["Subtitle"].font.color.rgb = RGBColor(75, 85, 99)
    styles["Subtitle"].font.italic = False
    styles["Heading 1"].font.size = Pt(12)
    styles["Heading 1"].paragraph_format.space_before = Pt(14)
    document.add_paragraph("가온랩  |  내부 운영 자료", style="Subtitle")
    document.add_paragraph(item["title"], style="Title")
    intro = document.add_paragraph(item["purpose"] + ". 이 문서는 해당 제도 또는 행사의 운영 기준을 확인하는 자료다.")
    intro.paragraph_format.space_after = Pt(12)
    table = document.add_table(rows=3, cols=2)
    table.style = "Table Grid"
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), "D9D9D9")
        borders.append(element)
    tbl_pr.append(borders)
    for row, pair in zip(table.rows, [("문서번호", "GL-" + item["id"].upper()), ("작성일", item["date"]), ("버전 및 상태", item["version"] + " · " + item["status"])]):
        for cell, value in zip(row.cells, pair):
            cell.text = value
            if cell == row.cells[0]:
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), "F1F5F9")
                cell._tc.get_or_add_tcPr().append(shade)
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.name = "Malgun Gothic"
                    run.font.size = Pt(8.5)
    for heading, paragraphs in item["sections"]:
        document.add_heading(heading, level=1)
        for paragraph in paragraphs:
            document.add_paragraph(paragraph)
    footer = section.footer.paragraphs[0]
    footer.text = "가온랩 피플운영·컬처팀  |  테스트용 가상 자료"
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in footer.runs:
        run.font.name = "Malgun Gothic"
        run.font.size = Pt(8)
    document.save(target)


def make_pdf(item, target):
    styles = {
        "eyebrow": ParagraphStyle("eyebrow", fontName="Malgun", fontSize=8, leading=13, textColor=colors.HexColor("#4B5563")),
        "title": ParagraphStyle("title", fontName="Malgun", fontSize=18, leading=27, spaceAfter=13),
        "intro": ParagraphStyle("intro", fontName="Malgun", fontSize=9.5, leading=17, spaceAfter=13),
        "heading": ParagraphStyle("heading", fontName="Malgun", fontSize=11, leading=17, spaceBefore=13, spaceAfter=6),
        "body": ParagraphStyle("body", fontName="Malgun", fontSize=9.2, leading=16, spaceAfter=7),
        "meta": ParagraphStyle("meta", fontName="Malgun", fontSize=8.5, leading=13),
    }
    story = [Paragraph("가온랩  |  내부 운영 자료", styles["eyebrow"]),
             Paragraph(item["title"], styles["title"]),
             Paragraph(item["purpose"] + ". 이 문서는 해당 제도 또는 행사의 운영 기준을 확인하는 자료다.", styles["intro"])]
    meta = [[Paragraph("문서번호", styles["meta"]), Paragraph("GL-" + item["id"].upper(), styles["meta"])],
            [Paragraph("작성일", styles["meta"]), Paragraph(item["date"], styles["meta"])],
            [Paragraph("버전 및 상태", styles["meta"]), Paragraph(item["version"] + " · " + item["status"], styles["meta"])]]
    table = Table(meta, colWidths=[100, 340], hAlign="LEFT")
    table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), .5, colors.HexColor("#D9D9D9")),
                               ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#F1F5F9")),
                               ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                               ("LEFTPADDING", (0, 0), (-1, -1), 9), ("RIGHTPADDING", (0, 0), (-1, -1), 9),
                               ("TOPPADDING", (0, 0), (-1, -1), 7), ("BOTTOMPADDING", (0, 0), (-1, -1), 7)]))
    story += [table, Spacer(1, 8)]
    for heading, paragraphs in item["sections"]:
        block = [Paragraph(heading, styles["heading"])]
        block += [Paragraph(p.replace("&", "&amp;").replace("<", "&lt;"), styles["body"]) for p in paragraphs]
        story.append(KeepTogether(block))
    def footer(canvas, pdf):
        canvas.setFont("Malgun", 7.5)
        canvas.setFillColor(colors.HexColor("#6B7280"))
        canvas.drawRightString(552, 35, f"가온랩 피플운영·컬처팀  |  테스트용 가상 자료  |  {pdf.page}쪽")
    SimpleDocTemplate(str(target), pagesize=(595.28, 841.89), leftMargin=74, rightMargin=74,
                      topMargin=67, bottomMargin=65).build(story, onFirstPage=footer, onLaterPages=footer)


for item in DOCS:
    (TEXT_DIR / f"{item['id']}.txt").write_text(source_text(item), encoding="utf-8")
    target = DOC_DIR / f"{item['id']}.{item['ext']}"
    (make_pdf if item["ext"] == "pdf" else make_docx)(item, target)

family = ["김", "이", "박", "최", "정", "강", "조", "윤"]
given = ["서연", "지훈", "민지", "도윤", "하은", "시우", "유진", "현우", "지민", "수빈", "예준", "서윤", "주원", "다은", "태윤", "가은", "준서", "채원", "민준", "지아"]
teams = ["마케팅팀", "개발팀", "운영팀", "디자인팀"]
employees = [dict(id=i + 1, employeeNo=f"GL-{i+1:04d}", name=family[i // 20] + given[i % 20],
                  team=teams[i % 4], email=f"employee{i+1:03d}@gaonlab.example", employmentStatus="재직")
             for i in range(160)]

def records(project, applied_ids, confirmed_ids=None, cancelled_ids=()):
    confirmed = set(confirmed_ids or ())
    return [dict(projectId=project, employeeId=i, status="confirmed" if i in confirmed else "applied",
                 checkedAt="2026-10-06T10:00") for i in applied_ids] + [
        dict(projectId=project, employeeId=i, status="cancelled", checkedAt="2026-10-06T10:00") for i in cancelled_ids]

projects = [
    dict(id="health-2026", targetIds=list(range(1, 121)), requiredIds=list(range(1, 31)), lastCheckedAt="2026-10-06T10:00"),
    dict(id="learning-2026q4", targetIds=list(range(41, 121)), requiredIds=[], lastCheckedAt="2026-10-06T10:00"),
    dict(id="workshop-2026", targetIds=list(range(1, 101)), requiredIds=list(range(1, 13)), lastCheckedAt="2026-10-06T10:00"),
    dict(id="photo-2026", targetIds=list(range(81, 161)), requiredIds=[], lastCheckedAt="2026-10-06T10:00"),
]
applications = records("health-2026", list(range(1, 26)) + list(range(31, 114)), list(range(1, 21)), [114])
applications += records("learning-2026q4", list(range(41, 71)), list(range(41, 51)), [71])
applications += records("workshop-2026", list(range(1, 10)) + list(range(13, 55)), list(range(1, 7)), [55])

(DATA_DIR / "employees.json").write_text(json.dumps(employees, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
(DATA_DIR / "projects.json").write_text(json.dumps(projects, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
(DATA_DIR / "applications.json").write_text(json.dumps(applications, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
(DATA_DIR / "cases.json").write_text(json.dumps(CASES, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
(DATA_DIR / "documents.json").write_text(json.dumps([{k: v for k, v in d.items() if k != "sections"} for d in DOCS], ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
print(f"Built {len(DOCS)} documents for {len(CASES)} cases and {len(employees)} fictional employees")
